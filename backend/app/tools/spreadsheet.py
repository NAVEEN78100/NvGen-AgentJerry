"""Spreadsheet reading and column-role inference (CSV + XLSX)."""
from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[A-Za-z]{2,}$")

# Semantic roles -> header synonyms (normalised, lowercase, no punctuation)
ROLE_SYNONYMS: dict[str, list[str]] = {
    "email": ["email", "emailaddress", "mail", "contactemail", "workemail", "emailid"],
    "name": ["name", "fullname", "contactname", "leadname", "candidatename", "customername", "clientname", "contact"],
    "first_name": ["firstname", "first", "givenname"],
    "company": ["company", "companyname", "organization", "organisation", "account", "business", "client"],
    "industry": ["industry", "sector", "vertical", "segment"],
    "status": ["status", "leadstatus", "stage", "pipelinestage", "state", "paymentstatus", "applicationstatus"],
    "amount": ["amount", "value", "dealvalue", "total", "invoiceamount", "amountdue", "balance", "revenue"],
    "due_date": ["duedate", "due", "paymentdue", "deadline"],
    "role": ["role", "position", "jobtitle", "appliedfor", "title"],
    "interest": ["interest", "product", "productinterest", "interestedin", "need", "requirement"],
    "source": ["source", "leadsource", "channel", "event"],
    "notes": ["notes", "note", "comments", "remarks", "lastinteraction", "context"],
    "owner": ["owner", "salesrep", "rep", "accountowner", "assignedto"],
    "invoice": ["invoice", "invoiceno", "invoicenumber", "invoiceid", "reference"],
    "date": ["date", "createdon", "dateadded", "created", "leaddate", "applieddate"],
}


def _norm(h: str) -> str:
    return re.sub(r"[^a-z0-9]", "", h.lower())


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == datetime.min.time() else v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


class SpreadsheetError(ValueError):
    pass


def read_table(path: str | Path, kind: str | None = None) -> tuple[list[str], list[dict[str, str]]]:
    path = Path(path)
    kind = (kind or path.suffix.lstrip(".")).lower()
    if not path.exists():
        raise SpreadsheetError(f"File not found: {path.name}")
    if kind == "csv":
        raw = path.read_bytes()
        for enc in ("utf-8-sig", "utf-8", "latin-1"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text), dialect)
        matrix = [row for row in reader]
    elif kind in ("xlsx", "xlsm"):
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        matrix = [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
    else:
        raise SpreadsheetError("Only .csv and .xlsx spreadsheets are supported.")

    matrix = [r for r in matrix if any(_cell(c) for c in r)]
    if not matrix:
        raise SpreadsheetError("The spreadsheet is empty.")
    header = [_cell(h) or f"Column {i + 1}" for i, h in enumerate(matrix[0])]
    # de-duplicate header names
    seen: dict[str, int] = {}
    for i, h in enumerate(header):
        if h in seen:
            seen[h] += 1
            header[i] = f"{h} ({seen[h]})"
        else:
            seen[h] = 1
    rows = []
    for r in matrix[1:]:
        cells = [_cell(c) for c in r] + [""] * (len(header) - len(r))
        rows.append({header[i]: cells[i] for i in range(len(header))})
    return header, rows


def infer_roles(columns: list[str], rows: list[dict[str, str]] | None = None) -> dict[str, str]:
    """Map semantic roles (email, name, status, ...) to actual column names."""
    roles: dict[str, str] = {}
    normed = {c: _norm(c) for c in columns}
    for role, syns in ROLE_SYNONYMS.items():
        # exact match first, then contains
        for c, n in normed.items():
            if n in syns and c not in roles.values():
                roles[role] = c
                break
        if role not in roles:
            for c, n in normed.items():
                if any(s in n for s in syns if len(s) > 3) and c not in roles.values():
                    roles[role] = c
                    break
    # email by content if header didn't say so
    if "email" not in roles and rows:
        for c in columns:
            vals = [r.get(c, "") for r in rows[:20] if r.get(c)]
            if vals and sum(bool(EMAIL_RE.match(v)) for v in vals) / len(vals) > 0.8:
                roles["email"] = c
                break
    return roles


def distinct_values(rows: list[dict[str, str]], column: str, limit: int = 12) -> list[tuple[str, int]]:
    counts: dict[str, int] = {}
    for r in rows:
        v = (r.get(column) or "").strip()
        if v:
            counts[v] = counts.get(v, 0) + 1
    return sorted(counts.items(), key=lambda kv: -kv[1])[:limit]


def parse_number(v: str) -> float | None:
    if v is None:
        return None
    s = re.sub(r"[^0-9.\-]", "", str(v))
    try:
        return float(s) if s not in ("", "-", ".") else None
    except ValueError:
        return None


def parse_date(v: str) -> date | None:
    v = (v or "").strip()
    candidates = {v, v.split(" ")[0] if len(v) > 10 and v[4] == "-" else v}
    for cand in candidates:
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%d %b %Y", "%d %B %Y", "%Y/%m/%d"):
            try:
                return datetime.strptime(cand, fmt).date()
            except ValueError:
                continue
    try:
        return datetime.fromisoformat(v).date()
    except ValueError:
        return None


def first_name(full: str) -> str:
    full = (full or "").strip()
    if not full:
        return ""
    parts = full.split()
    if parts[0].rstrip(".").lower() in {"mr", "mrs", "ms", "dr", "prof"} and len(parts) > 1:
        return parts[1]
    return parts[0]


def write_csv(path: Path, columns: list[str], rows: list[dict[str, Any]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=columns, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow(r)
